from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.contrib import messages  # لإظهار رسائل
from django.db.models import Q
from accounting.models import Account   # ⭐ الحسابات
from .models import Product, BundleComponent, Category  # ✅ إضافة Category
from django.db.models.deletion import ProtectedError
from django.db import IntegrityError


from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.contrib import messages

from accounting.models import Account

from .models import (
    Product,
    ProductImage,
    BundleComponent,
    Category,
)

from django.db.models.deletion import ProtectedError
from django.db import IntegrityError


# ================================
#   قائمة المنتجات
# ================================
def products_list(request):

    q = request.GET.get("q", "")

    if not getattr(request, "company", None):
        products = Product.objects.none()
    else:
        products = Product.objects.filter(
            company=request.company
        )

    if q:
        products = products.filter(
            name__icontains=q
        )

    return render(
        request,
        "products/products_list.html",
        {
            "products": products
        }
    )


# ================================
#   إضافة منتج جديد
# ================================
def product_add(request):

    print("========== PRODUCT DEBUG ==========")
    print("USER:", request.user)
    print("COMPANY:", getattr(request, "company", None))

    if not getattr(request, "company", None):

        messages.error(
            request,
            "لا يمكن إضافة منتج قبل ربط المستخدم بشركة."
        )

        return redirect("/")

    products = Product.objects.filter(
        company=request.company
    )

    accounts = Account.objects.all().order_by("code")

    categories = Category.objects.filter(
        company=request.company,
        active=True
    )

    if request.method == "POST":

        # ==========================================
        # نوع المنتج
        # ==========================================

        product_type = request.POST.get(
            "type",
            "normal"
        )

        # ==========================================
        # التصنيف
        # ==========================================

        category_id = (
            request.POST.get("category")
            or None
        )

        category = None

        if category_id:

            category = Category.objects.filter(
                company=request.company,
                id=category_id
            ).first()

        # ==========================================
        # الصورة الرئيسية
        # ==========================================

        main_image = request.FILES.get("image")

        # ==========================================
        # إنشاء المنتج
        # ==========================================

        barcode = request.POST.get("barcode") or None

        if barcode:

            barcode_exists = Product.objects.filter(
                company=request.company,
                barcode=barcode,
            ).exists()

            if barcode_exists:

                messages.error(
                    request,
                    "هذا الباركود مستخدم بالفعل داخل هذه الشركة."
                )

                return render(
                    request,
                    "products/product_form.html",
                    {
                        "products": products,
                        "accounts": accounts,
                        "categories": categories,
                    }
                )


        product = Product.objects.create(

            company=request.company,

            name=request.POST.get("name"),

            sku=request.POST.get("sku"),

            barcode=barcode,
            category=category,

            purchase_price=(
                request.POST.get("purchase_price")
                or 0
            ),

            sale_price=(
                request.POST.get("sale_price")
                or 0
            ),

            min_stock=(
                request.POST.get("min_stock")
                or 0
            ),

            alert_stock=(
                request.POST.get("alert_stock")
                or 0
            ),

            inventory_account_id=(
                request.POST.get("inventory_account")
                or None
            ),

            cost_account_id=(
                request.POST.get("cost_account")
                or None
            ),

            revenue_account_id=(
                request.POST.get("revenue_account")
                or None
            ),

            description=request.POST.get(
                "description",
                ""
            ),

            image=main_image,

            type=product_type,

            active=True,
        )

        # ==========================================
        # الصور الإضافية المتعددة
        # ==========================================

        gallery_images = request.FILES.getlist(
            "gallery_images"
        )

        for index, image_file in enumerate(
            gallery_images
        ):

            ProductImage.objects.create(

                product=product,

                image=image_file,

                sort_order=index,

                is_main=False,

            )

        # ==========================================
        # مكونات المنتج المركب
        # ==========================================

        if product_type == "bundle":

            count = int(
                request.POST.get(
                    "bundle_count",
                    0
                )
            )

            for i in range(
                1,
                count + 1
            ):

                comp_id = request.POST.get(
                    f"component_{i}"
                )

                qty = request.POST.get(
                    f"qty_{i}"
                )

                if comp_id and qty:

                    comp = Product.objects.filter(
                        company=request.company,
                        id=comp_id
                    ).first()

                    if comp:

                        BundleComponent.objects.create(

                            product=product,

                            component=comp,

                            quantity=qty,

                        )

        # ==========================================
        # الرجوع للفاتورة
        # ==========================================

        if request.GET.get("return") == "invoice":

            return redirect(
                f"/sales/invoices/add/?product_id={product.id}"
            )

        return redirect(
            "/products/"
        )

    return render(
        request,
        "products/product_form.html",
        {
            "products": products,
            "accounts": accounts,
            "categories": categories,
        }
    )


# ================================
#   تعديل منتج
# ================================
def product_edit(request, pk):

    if not getattr(request, "company", None):

        messages.error(
            request,
            "لا يمكن تعديل منتج قبل ربط المستخدم بشركة."
        )

        return redirect("/")

    product = get_object_or_404(
        Product,
        pk=pk,
        company=request.company
    )

    products = Product.objects.filter(
        company=request.company
    ).exclude(
        id=pk
    )

    accounts = Account.objects.all().order_by(
        "code"
    )

    categories = Category.objects.filter(
        company=request.company,
        active=True
    )

    if request.method == "POST":

        # ==========================================
        # البيانات الأساسية
        # ==========================================

        product.name = request.POST.get(
            "name"
        )

        product.sku = request.POST.get(
            "sku"
        )

        barcode = request.POST.get(
            "barcode"
        ) or None

        if barcode:

            barcode_exists = Product.objects.filter(
                company=request.company,
                barcode=barcode,
            ).exclude(
                id=product.id
            ).exists()

            if barcode_exists:

                messages.error(
                    request,
                    "هذا الباركود مستخدم بالفعل داخل هذه الشركة."
                )

                return render(
                    request,
                    "products/product_form.html",
                    {
                        "product": product,
                        "products": products,
                        "components": BundleComponent.objects.filter(
                            product=product
                        ),
                        "accounts": accounts,
                        "categories": categories,
                        "gallery_images": ProductImage.objects.filter(
                            product=product
                        ).order_by(
                            "sort_order",
                            "id"
                        ),
                    }
                )

        product.barcode = barcode
        product.purchase_price = (
            request.POST.get(
                "purchase_price"
            )
            or 0
        )

        product.sale_price = (
            request.POST.get(
                "sale_price"
            )
            or 0
        )

        product.min_stock = (
            request.POST.get(
                "min_stock"
            )
            or 0
        )

        product.alert_stock = (
            request.POST.get(
                "alert_stock"
            )
            or 0
        )

        # ==========================================
        # الحسابات
        # ==========================================

        product.inventory_account_id = (
            request.POST.get(
                "inventory_account"
            )
            or None
        )

        product.cost_account_id = (
            request.POST.get(
                "cost_account"
            )
            or None
        )

        product.revenue_account_id = (
            request.POST.get(
                "revenue_account"
            )
            or None
        )

        # ==========================================
        # التصنيف
        # ==========================================

        category_id = (
            request.POST.get(
                "category"
            )
            or None
        )

        if category_id:

            product.category = Category.objects.filter(
                company=request.company,
                id=category_id
            ).first()

        else:

            product.category = None

        # ==========================================
        # الوصف
        # ==========================================

        product.description = request.POST.get(
            "description",
            ""
        )

        # ==========================================
        # نوع المنتج
        # ==========================================

        product_type = request.POST.get(
            "type",
            "normal"
        )

        product.type = product_type

        # ==========================================
        # الصورة الرئيسية
        # ==========================================

        main_image = request.FILES.get(
            "image"
        )

        if main_image:

            product.image = main_image

        # ==========================================
        # حفظ المنتج
        # ==========================================

        product.save()

        # ==========================================
        # حذف الصور الإضافية المحددة
        # ==========================================

        delete_images = request.POST.getlist(
            "delete_gallery_images"
        )

        if delete_images:

            ProductImage.objects.filter(
                product=product,
                id__in=delete_images
            ).delete()

        # ==========================================
        # إضافة صور جديدة متعددة
        # ==========================================

        gallery_images = request.FILES.getlist(
            "gallery_images"
        )

        if gallery_images:

            last_image = (
                ProductImage.objects.filter(
                    product=product
                )
                .order_by("-sort_order")
                .first()
            )

            if last_image:

                next_sort_order = (
                    last_image.sort_order + 1
                )

            else:

                next_sort_order = 0

            for index, image_file in enumerate(
                gallery_images
            ):

                ProductImage.objects.create(

                    product=product,

                    image=image_file,

                    sort_order=(
                        next_sort_order + index
                    ),

                    is_main=False,

                )

        # ==========================================
        # تحديث مكونات Bundle
        # ==========================================

        if product_type == "bundle":

            BundleComponent.objects.filter(
                product=product
            ).delete()

            count = int(
                request.POST.get(
                    "bundle_count",
                    0
                )
            )

            for i in range(
                1,
                count + 1
            ):

                comp_id = request.POST.get(
                    f"component_{i}"
                )

                qty = request.POST.get(
                    f"qty_{i}"
                )

                if comp_id and qty:

                    comp = Product.objects.filter(
                        company=request.company,
                        id=comp_id
                    ).first()

                    if comp:

                        BundleComponent.objects.create(

                            product=product,

                            component=comp,

                            quantity=qty,

                        )

        else:

            # لو تغير المنتج من Bundle
            # إلى عادي أو خدمة
            BundleComponent.objects.filter(
                product=product
            ).delete()

        # ==========================================
        # الرجوع
        # ==========================================

        next_url = request.GET.get(
            "next"
        )

        if next_url:

            return redirect(
                next_url
            )

        return redirect(
            "/products/"
        )

    # ==========================================
    # البيانات المطلوبة للقالب
    # ==========================================

    components = BundleComponent.objects.filter(
        product=product
    )

    gallery_images = ProductImage.objects.filter(
        product=product
    ).order_by(
        "sort_order",
        "id"
    )

    return render(
        request,
        "products/product_form.html",
        {
            "product": product,

            "products": products,

            "components": components,

            "accounts": accounts,

            "categories": categories,

            "gallery_images": gallery_images,
        }
    )

# ================================
#   حذف منتج
# ================================
def product_delete(request, pk):

    if not getattr(request, "company", None):
        messages.error(
            request,
            "لا يمكن حذف منتج قبل ربط المستخدم بشركة."
        )
        return redirect("/")

    product = get_object_or_404(
        Product,
        pk=pk,
        company=request.company
    )

    # -----------------------------------------
    # التأكد أن المنتج ليس مكوّناً لمنتج مركب
    # -----------------------------------------
    linked_components = BundleComponent.objects.filter(
        component=product
    ).exists()

    if linked_components:
        messages.error(
            request,
            "لا يمكن حذف المنتج لأنه مستخدم كمكوّن في منتج مركب."
        )
        return redirect("/products/")

    try:
        product.delete()

        messages.success(
            request,
            "تم حذف المنتج بنجاح."
        )

    except ProtectedError as e:
        print("========== PRODUCT DELETE PROTECTED ERROR ==========")
        print(e)

        messages.error(
            request,
            "لا يمكن حذف المنتج لأنه مرتبط بسجلات أخرى في النظام."
        )

    except IntegrityError as e:
        print("========== PRODUCT DELETE INTEGRITY ERROR ==========")
        print(e)

        messages.error(
            request,
            "لا يمكن حذف المنتج بسبب وجود سجلات مرتبطة به."
        )

    except Exception as e:
        print("========== PRODUCT DELETE ERROR ==========")
        print(type(e).__name__)
        print(e)

        messages.error(
            request,
            f"حدث خطأ أثناء حذف المنتج: {e}"
        )

    return redirect("/products/")

# ================================
#   عرض منتج
# ================================
def product_view(request, pk):

    if not getattr(request, "company", None):
        messages.error(request, "لا يمكن عرض منتج قبل ربط المستخدم بشركة.")
        return redirect("/")

    product = get_object_or_404(Product, pk=pk, company=request.company)
    components = BundleComponent.objects.filter(product=product)

    return render(request, "products/product_view.html", {
        "product": product,
        "components": components
    })


# ================================
#   بحث المنتجات (API للفواتير)
# ================================
def search_products(request):

    q = request.GET.get("q", "").strip()

    if not getattr(request, "company", None):

        return JsonResponse(
            [],
            safe=False
        )

    products = Product.objects.filter(
        company=request.company,
        active=True
    )

    if q:

        products = products.filter(
            Q(name__icontains=q)
            | Q(sku__icontains=q)
            | Q(barcode__icontains=q)
        )

    products = products[:20]

    results = [{
        "id": p.id,
        "name": p.name,
        "barcode": p.barcode,
        "price": float(p.sale_price or 0),
    } for p in products]

    return JsonResponse(
        results,
        safe=False
    )
# ==========================================
# تصدير قائمة المنتجات: Excel وWord وPDF
# ==========================================

def _products_export_queryset(request):
    """إرجاع منتجات الشركة المحددة فقط، مع احترام البحث بالاسم."""
    company = getattr(request, "company", None)

    if company is None:
        return None

    products = Product.objects.filter(company=company)

    q = request.GET.get("q", "").strip()
    if q:
        products = products.filter(name__icontains=q)

    return products.select_related("category").order_by("name", "id")


def _products_export_rows(products):
    """تجهيز أعمدة موحدة للملفات الثلاثة."""
    headers = [
        "اسم المنتج",
        "كود المنتج",
        "الباركود",
        "التصنيف",
        "سعر الشراء",
        "سعر البيع",
        "الكمية المتاحة",
    ]

    rows = []
    for product in products:
        rows.append([
            product.name or "",
            product.sku or "",
            product.barcode or "",
            product.category.name if product.category else "",
            float(product.purchase_price or 0),
            float(product.sale_price or 0),
            float(product.current_stock or 0),
        ])

    return headers, rows


def products_export_excel(request):
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    products = _products_export_queryset(request)
    if products is None:
        return HttpResponse("لا توجد شركة محددة للمستخدم.", status=403)

    headers, rows = _products_export_rows(products)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "المنتجات"
    sheet.sheet_view.rightToLeft = True

    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="217346")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center")

    for column_cells in sheet.columns:
        column_index = column_cells[0].column
        max_length = max(
            len(str(cell.value or "")) for cell in column_cells
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 3, 14), 35
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = 'attachment; filename="products.xlsx"'
    return response


def products_export_word(request):
    from io import BytesIO
    from django.http import HttpResponse
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    products = _products_export_queryset(request)
    if products is None:
        return HttpResponse("لا توجد شركة محددة للمستخدم.", status=403)

    headers, rows = _products_export_rows(products)

    document = Document()
    document.add_heading("قائمة المنتجات", level=1)

    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )
    response["Content-Disposition"] = 'attachment; filename="products.docx"'
    return response


def products_export_pdf(request):
    from io import BytesIO
    from pathlib import Path
    from django.conf import settings
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    import arabic_reshaper
    from bidi.algorithm import get_display

    products = _products_export_queryset(request)
    if products is None:
        return HttpResponse("لا توجد شركة محددة للمستخدم.", status=403)

    headers, rows = _products_export_rows(products)

    font_path = Path(settings.BASE_DIR) / "static" / "fonts" / "Amiri-Regular.ttf"
    if not font_path.exists():
        return HttpResponse(
            "تعذر إنشاء PDF: ملف الخط العربي Amiri-Regular.ttf غير موجود.",
            status=500,
            content_type="text/plain; charset=utf-8",
        )

    if "AmiriExport" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AmiriExport", str(font_path)))

    def shape_arabic(value):
        text = str(value if value is not None else "")
        return get_display(arabic_reshaper.reshape(text))

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ArabicExportTitle",
        parent=styles["Title"],
        fontName="AmiriExport",
        fontSize=18,
        leading=24,
        alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "ArabicExportCell",
        parent=styles["Normal"],
        fontName="AmiriExport",
        fontSize=8,
        leading=11,
        alignment=TA_RIGHT,
    )
    header_style = ParagraphStyle(
        "ArabicExportHeader",
        parent=cell_style,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    table_data = [
        [Paragraph(shape_arabic(header), header_style) for header in
         ["الكمية المتاحة", "سعر البيع", "سعر الشراء", "التصنيف",
          "الباركود", "كود المنتج", "اسم المنتج"]]
    ]

    # ترتيب الأعمدة من اليسار إلى اليمين، بحيث يظهر اسم المنتج أقصى اليمين.
    for row in rows:
        reversed_row = list(reversed(row))
        table_data.append([
            Paragraph(shape_arabic(value), cell_style)
            for value in reversed_row
        ])

    table = Table(
        table_data,
        repeatRows=1,
        colWidths=[65, 65, 65, 90, 95, 80, 230],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B42318")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))

    story = [
        Paragraph(shape_arabic("قائمة المنتجات"), title_style),
        Spacer(1, 12),
        table,
    ]
    document.build(story)
    output.seek(0)

    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="products.pdf"'
    return response